"""Tests for tailor.py: rec parsing, budgets, PPTX edit application."""
import pytest
from pptx import Presentation
from pptx.util import Inches

from job_agent.tailor import (
    parse_recommendation,
    budgets_from_markdown,
    apply_edits,
    format_report,
    MOVE_FIRST,
)

REC_MD = """**Resume version:** ai_focused
**Why:** AI-focused is the default for this search.

**Work Experience | Acme Corp**

- **Original:** Led migration of batch pipelines to streaming.
- **Suggested:** Led migration of batch pipelines to real-time streaming.
- **Why:** "Real-time" matches the JD language.

**Work Experience | Acme Corp**

- **Move:** Third bullet about cost savings.
- **After:** First bullet about data pipelines.
- **Why:** Cost impact should lead.
"""


def test_parse_recommendation():
    version, edits = parse_recommendation(REC_MD)
    assert version == "ai_focused"
    assert len(edits) == 2
    assert edits[0]["operation"] == "replace"
    assert edits[0]["original"] == "Led migration of batch pipelines to streaming."
    assert edits[0]["suggested"] == "Led migration of batch pipelines to real-time streaming."
    assert edits[1]["operation"] == "reorder"
    assert edits[1]["original"] == "Third bullet about cost savings."
    assert edits[1]["suggested"] == "First bullet about data pipelines."


def test_budgets_from_markdown():
    md = """# Resume

## Introduction

Data leader with ten years of experience.

## Work Experience

### Acme Corp | Chicago, IL | 2020-2024

- Led migration of batch pipelines.
- Built a team of five.
"""
    budgets = budgets_from_markdown(md)
    by_loc = {b["locator"]: b for b in budgets}
    assert by_loc["Introduction"]["budget"] == len("Data leader with ten years of experience.")
    assert by_loc["Work Experience | Acme Corp bullet 1"]["budget"] == len("Led migration of batch pipelines.")
    assert by_loc["Work Experience | Acme Corp bullet 2"]["budget"] == len("Built a team of five.")


def _make_pptx(path, bullets):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    # reuse the default first paragraph for the first bullet
    tf.paragraphs[0].text = bullets[0]
    for b in bullets[1:]:
        tf.add_paragraph().text = b
    prs.save(str(path))
    return path


def _bullets_of(pptx_path):
    prs = Presentation(str(pptx_path))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    return out


def test_apply_replace(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["Alpha bullet", "Beta bullet"])
    out = tmp_path / "out.pptx"
    edits = [{"operation": "replace", "section": "S",
              "original": "Alpha bullet", "suggested": "Alpha revised", "reason": "x"}]
    report = apply_edits(src, edits, out)
    assert len(report["applied"]) == 1
    assert report["skipped"] == []
    assert _bullets_of(out) == ["Alpha revised", "Beta bullet"]


def test_apply_replace_no_match(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["Alpha bullet"])
    out = tmp_path / "out.pptx"
    edits = [{"operation": "replace", "section": "S",
              "original": "Something entirely different", "suggested": "New", "reason": "x"}]
    report = apply_edits(src, edits, out)
    assert report["applied"] == []
    assert len(report["skipped"]) == 1
    assert _bullets_of(out) == ["Alpha bullet"]  # unchanged


def test_apply_reorder(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["First", "Second", "Third"])
    out = tmp_path / "out.pptx"
    edits = [{"operation": "reorder", "section": "S",
              "original": "Third", "suggested": "First", "reason": "x"}]
    report = apply_edits(src, edits, out)
    assert len(report["applied"]) == 1
    assert _bullets_of(out) == ["First", "Third", "Second"]


def test_apply_reorder_to_top(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["First", "Second", "Third"])
    out = tmp_path / "out.pptx"
    edits = [{"operation": "reorder", "section": "S",
              "original": "Third", "suggested": MOVE_FIRST, "reason": "x"}]
    report = apply_edits(src, edits, out)
    assert len(report["applied"]) == 1
    assert _bullets_of(out) == ["Third", "First", "Second"]


def test_overflow_flagged(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["Short"])
    out = tmp_path / "out.pptx"
    long_text = "A much longer replacement that far exceeds the original bullet length"
    edits = [{"operation": "replace", "section": "S",
              "original": "Short", "suggested": long_text, "reason": "x"}]
    report = apply_edits(src, edits, out)
    assert len(report["flagged"]) >= 1
    assert "verify layout" in report["flagged"][0]["detail"]


def test_dry_run_writes_nothing(tmp_path):
    src = _make_pptx(tmp_path / "src.pptx", ["Alpha"])
    out = tmp_path / "out.pptx"
    edits = [{"operation": "replace", "section": "S",
              "original": "Alpha", "suggested": "Beta", "reason": "x"}]
    report = apply_edits(src, edits, out, dry_run=True)
    assert report["dry_run"] is True
    assert not out.exists()
    assert len(report["applied"]) == 1


def test_format_report():
    report = {
        "output_path": "out.pptx", "source_pptx": "src.pptx",
        "version": "ai_focused", "edits_found": 2, "dry_run": False,
        "applied": [{"operation": "replace", "section": "S", "detail": "ok"}],
        "flagged": [{"operation": "replace", "section": "S2", "detail": "check it"}],
        "skipped": [],
    }
    text = format_report(report)
    assert "out.pptx" in text
    assert "Needs your eye" in text


# ---------------------------------------------------------------------------
# Drive flow (mocked services)
# ---------------------------------------------------------------------------

class _FakeFiles:
    def __init__(self, existing_folders=None):
        self.existing = existing_folders or {}
        self.created = []
        self._resp = {}

    def list(self, q=None, fields=None, pageSize=None):
        self._q = q
        # extract name='...' from query
        import re as _re
        m = _re.search(r"name='([^']+)'", q or "")
        name = m.group(1) if m else ""
        self._resp = {"files": [{"id": self.existing[name]}] if name in self.existing else []}
        return self

    def create(self, body=None, media_body=None, fields=None):
        self.created.append(body)
        self._resp = {"id": "new-file-id", "webViewLink": "https://drive/fake"}
        return self

    def execute(self):
        return self._resp


class _FakeDrive:
    def __init__(self, existing_folders=None):
        self._files = _FakeFiles(existing_folders)

    def files(self):
        return self._files


def test_find_or_create_folder_existing():
    from job_agent.tailor import _find_or_create_folder
    drive = _FakeDrive(existing_folders={"2026_10": "folder-123"})
    fid = _find_or_create_folder(drive, "2026_10", "parent-1")
    assert fid == "folder-123"
    assert drive.files().created == []


def test_find_or_create_folder_creates():
    from job_agent.tailor import _find_or_create_folder
    drive = _FakeDrive()
    fid = _find_or_create_folder(drive, "2026_10", "parent-1")
    assert fid == "new-file-id"
    body = drive.files().created[0]
    assert body["name"] == "2026_10"
    assert body["parents"] == ["parent-1"]
    assert "folder" in body["mimeType"]


def test_upload_tailored_pptx(monkeypatch, tmp_path):
    import json as _json
    from job_agent import tailor as tailor_mod
    drive = _FakeDrive(existing_folders={"2026_10": "folder-123"})
    # upload_tailored_pptx imports get_google_services from .drive_sync inside
    # the function, so patch it there.
    import job_agent.drive_sync as ds
    monkeypatch.setattr(ds, "get_google_services", lambda profile: (drive, None))

    out_dir = tmp_path / "app"
    out_dir.mkdir()
    (out_dir / "score.json").write_text(_json.dumps({"company": "Acme Corp"}))
    pptx = tmp_path / "t.pptx"
    pptx.write_bytes(b"fake")
    profile = {"tailoring": {"output_parent_folder_id": "parent-1"},
               "profile_sync": {"credentials_file": "x"}}
    result = tailor_mod.upload_tailored_pptx(profile, pptx, out_dir)
    assert result["folder_id"] == "folder-123"
    assert result["filename"].startswith("TheoKelly_CV_acme-corp_")
    assert result["filename"].endswith(".pptx")
    body = drive.files().created[0]
    assert body["parents"] == ["folder-123"]


def test_tailor_from_output_dir_drive_base(monkeypatch, tmp_path):
    from job_agent import tailor as tailor_mod
    import job_agent.drive_sync as ds

    # fake output dir with a rec
    out_dir = tmp_path / "app"
    out_dir.mkdir()
    (out_dir / "resume_recommendation.md").write_text(
        "**Resume version:** ai_focused\n\n**S**\n\n"
        "- **Original:** Alpha\n- **Suggested:** Beta\n- **Why:** x\n")
    src = tmp_path / "base.pptx"
    _make_pptx(src, ["Alpha", "Gamma"])

    def fake_download(profile, dest):
        import shutil as _sh
        _sh.copy(str(src), str(dest))
        return dest
    monkeypatch.setattr(tailor_mod, "download_base_pptx", fake_download)

    drive = _FakeDrive()
    monkeypatch.setattr(ds, "get_google_services", lambda profile: (drive, None))

    profile = {"tailoring": {"base_pptx_drive_id": "base-1",
                              "output_parent_folder_id": "parent-1",
                              "upload_to_drive": True},
               "profile_sync": {"credentials_file": "x"}}
    report = tailor_mod.tailor_from_output_dir(out_dir, profile)
    assert "Drive base (base-1)" in report["source_pptx"]
    assert len(report["applied"]) == 1
    assert (out_dir / "resume_tailored.pptx").exists()
    assert report["drive_upload"]["filename"].endswith(".pptx")
